import unittest, math
import engine as E, verify, narrative

P = {"income": 25e6, "expenses": 12e6, "investments": 50e6, "assets": 120e6, "debt": 80e6, "debt_payment": 3e6, "years": 10, "liquidity": "medium", "risk": "medium", "debt_rate": 8}


class Independent(unittest.TestCase):
    def test_reference_checks(self):
        for seed in (7, 11, 23):
            good, total, fails = verify.run_checks(seed, 40)
            self.assertEqual(fails, [], fails[:5])
            self.assertEqual(good, total)


class EdgeCases(unittest.TestCase):
    def test_all_zero_and_garbage(self):
        for raw in ({}, None, {"income": 0}, {"income": "abc", "expenses": -5}, {"income": float("nan"), "debt": float("inf")}, {"income": 1e30}):
            a = E.analysis(raw)
            self.assertEqual(len(a["scenarios"]), 5)
            E.whatif(raw, {"income": {"sign": "-", "pct": 100}})
            E.goals_analyze(raw, [{"name": "x", "target": 10, "current": 0, "years": 1}])

    def test_negative_cashflow_draws_down(self):
        a = E.analysis(dict(P, income=8e6))
        base = next(s for s in a["scenarios"] if s["key"] == "baseline")
        self.assertLess(base["net_cash_flow"], 0)
        self.assertGreater(a["scenarios"][0]["shortfall_months"] + base["shortfall_months"], 0)
        self.assertTrue(all(math.isfinite(x) for x in base["net_worth_by_year"]))

    def test_payment_below_interest_flagged(self):
        issues = E.validate_profile(dict(P, debt_payment=100000, debt_rate=24))
        self.assertTrue(any("interest" in i["msg"] for i in issues))
        pl = E.goal_plan("Debt", 80e6, 0, 3, 0, 0, "debt", 24, 100000)
        self.assertTrue(pl["negative_amortization"])

    def test_validation_messages(self):
        msgs = " ".join(i["msg"] for i in E.validate_profile({"income": 0, "debt": 5e6, "debt_payment": 0}))
        self.assertIn("income is missing", msgs.lower()); self.assertIn("no monthly payment", msgs)
        self.assertTrue(any(i["level"] == "error" for i in E.validate_profile({"income": 5e6, "debt": 1, "debt_payment": 9e6})))


class GoalModel(unittest.TestCase):
    def test_no_double_counting(self):
        p = E.clean_profile(P)
        goals = E.default_goals(p)
        self.assertLessEqual(sum(g["current"] for g in goals), p["assets"] + p["investments"])
        bad = [dict(g, current=p["assets"] + p["investments"]) for g in goals[:2]]
        r = E.goals_analyze(P, bad)
        self.assertIn("errors", r); self.assertIn("same money", r["errors"][0])

    def test_capacity_shared_once(self):
        r = E.goals_analyze(dict(P, income=14e6))        # small capacity
        al = r["allocation"]
        self.assertLessEqual(al["total_allocated"], al["capacity"] + 1e-6)
        self.assertAlmostEqual(al["total_allocated"] + al["unallocated"], al["capacity"], 4)

    def test_debt_goal_model(self):
        rd = 0.12 / 12
        pay = E.pmt(100e6, rd, 36)
        pl = E.goal_plan("Debt", 100e6, 0, 3, 0, 0, "debt", 12, pay)
        self.assertEqual(pl["required_monthly"], 0)
        self.assertAlmostEqual(pl["completion"], 1.0, 6)
        self.assertLessEqual(abs(pl["months_at_rate"] - 36), 1)
        pl2 = E.goal_plan("Debt", 100e6, 0, 3, 0, 0, "debt", 12, pay * 0.8)
        self.assertGreater(pl2["required_monthly"], 0); self.assertLess(pl2["completion"], 1)

    def test_required_savings_consistent(self):
        pl = E.goal_plan("G", 500e6, 50e6, 5, 5e6, 0.06)
        self.assertEqual(pl["on_track"], pl["monthly_save"] >= pl["required_monthly"] - 0.5)
        ok = E.goal_plan("G", 500e6, 50e6, 5, pl["required_monthly"], 0.06)
        self.assertAlmostEqual(ok["completion"], 1.0, 6)
        self.assertLessEqual(ok["months_at_rate"], 60)
        self.assertAlmostEqual(ok["milestones"][-1]["balance"], 500e6, 0)

    def test_secondary_goal_and_assumptions(self):
        r = E.goals_analyze(P, None, None, {"income": 40e6, "save_rate": 90, "ret_sign": "-", "ret_pct": 2, "years": 8})
        a = r["assumed"]
        self.assertLessEqual(a["monthly_save"], max(a["net_cash_flow"], 0) + 1e-6)     # cannot save money that does not exist
        self.assertGreaterEqual(a["exceeds_cash_flow"], 0)
        self.assertIsNotNone(r["secondary"]); self.assertNotEqual(r["secondary"]["name"], r["goals"][r["primary"]]["name"])
        one = E.goals_analyze(P, [{"name": "Car", "target": 300e6, "current": 0, "years": 5}], None, {"income": 30e6, "save_rate": 30})
        self.assertTrue(one["secondary"]["suggested"])


class WhatIf(unittest.TestCase):
    def test_zero_change_equals_baseline(self):
        w = E.whatif(P, {})
        for r in w["rows"]:
            self.assertTrue(r["variance"] in (None, 0) or abs(r["variance"]) < 1e-6, r)

    def test_each_control_moves_results(self):
        base = E.whatif(P, {})
        d = E.whatif(P, {"debt": {"sign": "+", "pct": 30}})       # faster repayment: less debt left, earlier Debt Freedom
        self.assertLess(next(r for r in d["rows"] if r["label"].startswith("Debt Freedom"))["adjusted"], next(r for r in d["rows"] if r["label"].startswith("Debt Freedom"))["baseline"])
        for key, sign in (("income", "-"), ("expense", "+"), ("unforeseen", "+"), ("savings", "-")):
            w = E.whatif(P, {key: {"sign": sign, "pct": 30}})
            nw = next(r for r in w["rows"] if r["label"].startswith("Net worth"))
            self.assertLess(nw["adjusted"], nw["baseline"], key)
        up = E.whatif(P, {"income": {"sign": "+", "pct": 30}})
        self.assertGreater(next(r for r in up["rows"] if r["label"].startswith("Net worth"))["adjusted"], base["rows"][10]["adjusted"])

    def test_savings_rate_cap(self):
        w = E.whatif(P, {"savings": {"sign": "-", "pct": 100}})
        self.assertAlmostEqual(w["planned"], 0)
        pool = next(r for r in w["rows"] if r["label"].startswith("Investment pool"))
        self.assertLess(pool["adjusted"], pool["baseline"])


class TextAndGrounding(unittest.TestCase):
    def test_text_numbers(self):
        n = E.text_numbers("I earn 25 million a month, spend 12 triệu, debt 1.5 tỷ and 500k savings, 25,000,000")
        for v in (25e6, 12e6, 1.5e9, 500e3):
            self.assertIn(v, n)
        ok, rej = E.verify_against_text({"income": 25e6, "assets": 7e8, "risk": "high"}, "I earn 25 million, my risk is high")
        self.assertEqual(rej, ["assets"])

    def test_grounded_rate(self):
        a = E.analysis(P)
        ok, c = E.grounded_rate("Income is 25,000,000 VND and savings 10,000,000 VND, fake 99,999,999 VND", a)
        self.assertEqual((ok, c), (2, 3))

    def test_narrative_is_grounded(self):
        a = E.analysis(P); n = narrative.analysis(a)
        txt = "\n".join([n["table"], n["chart"], n["guidance"], n["recommendations"]] + list(n["scenarios"].values()))
        ok, c = E.grounded_rate(txt, a)
        self.assertGreaterEqual(ok / c, 0.95)


class ExcelParsing(unittest.TestCase):
    def test_template_roundtrip(self):
        r = E.parse_xlsx(E.template_xlsx())
        self.assertEqual(r["data"]["income"], 25e6); self.assertEqual(r["data"]["risk"], "medium"); self.assertEqual(r["missing"], [])

    def _wb(self, rows):
        from openpyxl import Workbook; import io
        wb = Workbook(); ws = wb.active
        for r in rows: ws.append(r)
        b = io.BytesIO(); wb.save(b); return b.getvalue()

    def test_bad_inputs(self):
        for raw in (b"", b"hello", b"PK\x03\x04garbage", b"PK" + b"\x00" * 100, b"x" * (E.MAX_XLSX + 1)):
            with self.assertRaises(ValueError): E.parse_xlsx(raw)
        with self.assertRaises(ValueError): E.parse_xlsx(self._wb([]))

    def test_messy_content(self):
        r = E.parse_xlsx(self._wb([("Monthly income", "25,000,000 VND"), ("Monthly expenses", -5), ("Assets", "abc"), ("Risk level", "extreme"),
                                   ("Time horizon", 99), ("Monthly income", 1), ("Total debt", "#REF!"), ("Lương tháng", 30000000), ("Thu nhập hàng tháng", "30.000.000")]))
        self.assertEqual(r["data"].get("income"), 25e6)
        self.assertNotIn("expenses", r["data"]); self.assertNotIn("assets", r["data"]); self.assertNotIn("risk", r["data"]); self.assertNotIn("years", r["data"])
        self.assertTrue(any("more than once" in w for w in r["warnings"]) and any("#REF" in w for w in r["warnings"]))
        self.assertEqual(r["unrecognized"][0]["label"], "Lương tháng")

    def test_exports_open_and_are_safe(self):
        from openpyxl import load_workbook; import io
        wb = load_workbook(io.BytesIO(E.report_xlsx(P)))
        self.assertEqual(wb.sheetnames, ["Financial Summary", "Health Assessment", "Scenarios", "Net Worth Projection", "Assumptions"])
        goals = [{"name": "=cmd|' /c calc'!A0", "target": 1e8, "current": 0, "years": 3}]
        wb = load_workbook(io.BytesIO(E.plan_xlsx(P, goals, None, {"income": 30e6})))
        self.assertIn("Funding Allocation", wb.sheetnames)
        self.assertTrue(str(wb["Goals"]["A2"].value).startswith("'"))
        with self.assertRaises(ValueError): E.plan_xlsx(P, [{"name": "x", "target": 0}], None, None)


if __name__ == "__main__":
    unittest.main()
